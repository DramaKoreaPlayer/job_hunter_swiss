"""Export des résultats vers Excel avec mise en forme."""

from __future__ import annotations

import os
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from rich.console import Console

import config
from scrapers.base import JobOffer

console = Console()


def _score_color(score: int) -> PatternFill:
    """Couleur de fond en fonction du score."""
    if score >= 75:
        return PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # Vert
    elif score >= 50:
        return PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")  # Jaune
    elif score >= 25:
        return PatternFill(start_color="FCD5B4", end_color="FCD5B4", fill_type="solid")  # Orange
    else:
        return PatternFill(start_color="E6E6E6", end_color="E6E6E6", fill_type="solid")  # Gris


def _prob_color(prob: int) -> PatternFill:
    """Couleur pour la probabilité de réussite."""
    if prob >= 70:
        return PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")  # Vert vif
    elif prob >= 45:
        return PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")  # Orange
    elif prob >= 20:
        return PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")  # Rouge clair
    else:
        return PatternFill(start_color="C0C0C0", end_color="C0C0C0", fill_type="solid")  # Gris


def export_to_excel(offers: list[JobOffer], search_queries: list[str]) -> str:
    """Exporte les offres vers un fichier Excel formaté.

    Returns:
        Chemin absolu du fichier créé.
    """
    # Créer le dossier results
    results_dir = os.path.join(os.path.dirname(__file__), "..", config.RESULTS_DIR)
    os.makedirs(results_dir, exist_ok=True)

    filename = config.EXCEL_FILENAME_TEMPLATE.format(
        date=date.today().strftime("%Y-%m-%d"),
        time=datetime.now().strftime("%Hh%M"),
    )
    filepath = os.path.join(results_dir, filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "Offres d'emploi"

    # ═══════════════════════════════════════════════
    # STYLES
    # ═══════════════════════════════════════════════
    title_font = Font(bold=True, color="FFFFFF", size=14)
    title_fill = PatternFill(start_color="1A1A2E", end_color="1A1A2E", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="16213E", end_color="16213E", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color="CCCCCC"),
        right=Side(style='thin', color="CCCCCC"),
        top=Side(style='thin', color="CCCCCC"),
        bottom=Side(style='thin', color="CCCCCC"),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    subtitle_font = Font(italic=True, color="666666", size=10)

    # ═══════════════════════════════════════════════
    # TITRE
    # ═══════════════════════════════════════════════
    headers = [
        "🏆 Score",
        "📊 Probabilité",
        "💼 Titre du poste",
        "🏢 Entreprise",
        "📍 Lieu",
        "⏰ Taux",
        "📄 Contrat",
        "🌐 Source",
        "🤖 Résumé IA",
        "🔗 Lien",
        "📋 Statut",
        "📝 Notes",
    ]

    # Ligne titre
    ws.merge_cells(f'A1:{get_column_letter(len(headers))}1')
    ws['A1'] = f"🔍 RECHERCHE D'EMPLOI — {config.DEFAULT_LOCATION} — {date.today().strftime('%d/%m/%Y')}"
    ws['A1'].font = title_font
    ws['A1'].fill = title_fill
    ws['A1'].alignment = center
    ws.row_dimensions[1].height = 40

    # Ligne sous-titre
    ws.merge_cells(f'A2:{get_column_letter(len(headers))}2')
    queries_display = ", ".join(search_queries[:5])
    if len(search_queries) > 5:
        queries_display += f" (+{len(search_queries) - 5} autres)"
    ws['A2'] = f"Recherches : {queries_display} | {len(offers)} offres trouvées"
    ws['A2'].font = subtitle_font
    ws['A2'].alignment = center
    ws.row_dimensions[2].height = 25

    # En-têtes
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin_border
    ws.row_dimensions[3].height = 35

    # ═══════════════════════════════════════════════
    # DONNÉES
    # ═══════════════════════════════════════════════
    for row_idx, offer in enumerate(offers, 4):
        row_data = [
            f"{offer.relevance_score}/100",
            f"{offer.success_probability}%",
            offer.title,
            offer.company,
            offer.location,
            offer.work_rate or "—",
            offer.contract_type or "—",
            offer.source,
            offer.ai_summary or "—",
            offer.url,
            "",  # Statut (à remplir manuellement)
            "",  # Notes
        ]

        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.alignment = left if col >= 3 else center
            cell.border = thin_border

            # Colorer les scores
            if col == 1:  # Score de pertinence
                cell.fill = _score_color(offer.relevance_score)
                cell.alignment = center
                cell.font = Font(bold=True, size=11)
            elif col == 2:  # Probabilité de réussite
                cell.fill = _prob_color(offer.success_probability)
                cell.alignment = center
                cell.font = Font(bold=True, size=11)
            elif col == 10:  # Lien
                cell.font = Font(color="0563C1", underline="single")

        # Hauteur de ligne adaptée
        ws.row_dimensions[row_idx].height = 50

    # ═══════════════════════════════════════════════
    # LARGEURS DE COLONNES
    # ═══════════════════════════════════════════════
    column_widths = [12, 14, 40, 28, 22, 12, 16, 20, 45, 50, 18, 28]
    for col, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    # ═══════════════════════════════════════════════
    # VALIDATION (liste déroulante pour le statut)
    # ═══════════════════════════════════════════════
    dv = DataValidation(
        type="list",
        formula1='"À postuler,Postulé,En attente,Entretien,Accepté,Refusé,Ignoré"',
        allow_blank=True,
    )
    dv.prompt = "Choisir le statut"
    dv.promptTitle = "Statut"
    ws.add_data_validation(dv)
    last_row = len(offers) + 3
    dv.add(f'K4:K{last_row}')

    # Figer les en-têtes
    ws.freeze_panes = 'A4'

    # Auto-filtre
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{last_row}"

    # ═══════════════════════════════════════════════
    # FEUILLE STATISTIQUES
    # ═══════════════════════════════════════════════
    ws_stats = wb.create_sheet("Statistiques")

    stat_data = [
        ("📊 Statistique", "Valeur"),
        ("Date de la recherche", date.today().strftime("%d/%m/%Y")),
        ("Région", config.DEFAULT_LOCATION),
        ("Nombre total d'offres", len(offers)),
        ("Offres haute pertinence (≥75)", sum(1 for o in offers if o.relevance_score >= 75)),
        ("Offres moyenne pertinence (50-74)", sum(1 for o in offers if 50 <= o.relevance_score < 75)),
        ("Offres basse pertinence (<50)", sum(1 for o in offers if o.relevance_score < 50)),
        ("Probabilité moyenne de réussite", f"{sum(o.success_probability for o in offers) / len(offers):.0f}%" if offers else "0%"),
        ("Sites utilisés", ", ".join(set(o.source.split(" + ")[0] for o in offers))),
        ("Termes de recherche", ", ".join(search_queries)),
    ]

    for row_idx, (label, value) in enumerate(stat_data, 1):
        ws_stats.cell(row=row_idx, column=1, value=label).font = Font(bold=(row_idx == 1))
        ws_stats.cell(row=row_idx, column=2, value=value)
        if row_idx == 1:
            ws_stats.cell(row=row_idx, column=1).fill = header_fill
            ws_stats.cell(row=row_idx, column=1).font = Font(bold=True, color="FFFFFF")
            ws_stats.cell(row=row_idx, column=2).fill = header_fill
            ws_stats.cell(row=row_idx, column=2).font = Font(bold=True, color="FFFFFF")

    ws_stats.column_dimensions['A'].width = 35
    ws_stats.column_dimensions['B'].width = 60

    # Sauvegarder (avec fallback si le fichier est déjà ouvert)
    try:
        wb.save(filepath)
    except PermissionError:
        # Fichier probablement ouvert dans Excel → ajouter un suffixe horodaté
        base, ext = os.path.splitext(filepath)
        timestamp = datetime.now().strftime("%H%M%S")
        filepath = f"{base}_{timestamp}{ext}"
        wb.save(filepath)
        console.print(
            "[yellow]⚠ Fichier précédent ouvert — sauvegardé sous un nouveau nom[/yellow]"
        )

    abs_path = os.path.abspath(filepath)
    console.print(f"\n[green]📁 Excel sauvegardé :[/green] [bold]{abs_path}[/bold]")

    return abs_path
